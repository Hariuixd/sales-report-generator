# sales_report_gui_merged.pyw
import os
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.pdfgen import canvas
import mysql.connector
import tkinter as tk
from tkinter import filedialog, messagebox
from datetime import datetime
from openpyxl import load_workbook

# --------- USER CONFIG ---------
# Update these with your MySQL credentials / DB name
DB_CONFIG = {
    "host": "localhost",
    "user": "root",
    "password": "Harixd006",   
    "database": "sales_db"   
}
# -------------------------------

# --- helper: find header row for Excel with stray rows ---
def find_header_row(file_path):
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        # look first 20 rows for a header-like row
        for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=20), start=0):
            values = [cell.value for cell in row if cell.value not in [None, '', ' ']]
            # if there are at least 3 unique values, treat as header
            if len(values) >= 3 and len(set(values)) == len(values):
                return i
    except Exception:
        pass
    return 0

# --- load and normalize data robustly ---
def load_data(path):
    try:
        ext = os.path.splitext(path)[-1].lower()
        if ext == ".csv":
            df = pd.read_csv(path)
        elif ext in [".xlsx", ".xls"]:
            header_row = find_header_row(path)
            df = pd.read_excel(path, engine="openpyxl", header=header_row)
        else:
            raise ValueError("Unsupported file format. Use .csv or .xlsx/.xls")

        # normalize column names
        df.columns = df.columns.astype(str).str.strip().str.lower()

        # automated column mapping (robust)
        column_map = {}
        for col in df.columns:
            if any(k in col for k in ["product", "item", "description", "goods"]):
                column_map[col] = "Product"
            elif any(k in col for k in ["qty", "units", "quantity", "nos", "rec qty", "rec_qty"]):
                column_map[col] = "Units Sold"
            elif any(k in col for k in ["rate", "price", "unit cost", "cost"]):
                column_map[col] = "Unit Price ($)"
            elif any(k in col for k in ["amount", "total", "sales", "value"]):
                column_map[col] = "Total Sales ($)"
            elif "date" in col:
                column_map[col] = "Date"

        df.rename(columns=column_map, inplace=True)

        # ensure required columns exist or compute them
        required = ["Product", "Units Sold", "Unit Price ($)"]
        for col in required:
            if col not in df.columns:
                raise ValueError(f"Missing column (expected or auto-mapped): {col}")

        # compute total if missing
        if "Total Sales ($)" not in df.columns:
            df["Total Sales ($)"] = pd.to_numeric(df["Units Sold"], errors="coerce") * pd.to_numeric(df["Unit Price ($)"], errors="coerce")

        # coerce types and drop bad rows
        df["Units Sold"] = pd.to_numeric(df["Units Sold"], errors="coerce")
        df["Unit Price ($)"] = pd.to_numeric(df["Unit Price ($)"], errors="coerce")
        df["Total Sales ($)"] = pd.to_numeric(df["Total Sales ($)"], errors="coerce")

        # date normalization if present
        if "Date" in df.columns:
            df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        else:
            df["Date"] = pd.NaT

        # drop rows missing crucial numeric info
        df = df.dropna(subset=["Units Sold", "Unit Price ($)"], how="any")
        return df
    except Exception as e:
        messagebox.showerror("Error", f"Failed to load data:\n{e}")
        return None

# --- report writers ---
def generate_word_report(summary, filename, include_chart):
    doc = Document()
    title = doc.add_heading("Sales Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.runs[0]
    run.font.size = Pt(20)
    run.bold = True

    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%d %B %Y')}")
    doc.add_paragraph("")

    doc.add_heading("Top Products Summary", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Product"
    hdr[1].text = "Units Sold"
    hdr[2].text = "Total Sales ($)"

    for _, row in summary.iterrows():
        r = table.add_row().cells
        r[0].text = str(row["Product"])
        r[1].text = f"{int(row['Units Sold']):,}"
        r[2].text = f"${row['Total Sales ($)']:.2f}"

    if include_chart and not summary.empty:
        chart_file = "temp_chart.png"
        try:
            chart = summary.set_index("Product")["Total Sales ($)"].sort_values(ascending=False).head(10)
            chart.plot(kind="bar", figsize=(6,4))
            plt.tight_layout()
            plt.savefig(chart_file, dpi=150)
            plt.close()
            doc.add_picture(chart_file, width=Inches(5.5))
        finally:
            if os.path.exists(chart_file):
                os.remove(chart_file)

    doc.save(filename)

def generate_pdf_report(summary, filename, include_chart):
    c = canvas.Canvas(filename)
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(300, 800, "Sales Report")
    c.setFont("Helvetica", 11)
    c.drawString(50, 780, f"Generated on: {datetime.now().strftime('%d %B %Y')}")
    y = 740
    c.drawString(50, y, "Product")
    c.drawString(300, y, "Units Sold")
    c.drawString(420, y, "Total Sales ($)")
    y -= 20

    for _, row in summary.iterrows():
        c.drawString(50, y, str(row["Product"])[:30])
        c.drawRightString(350, y, f"{int(row['Units Sold']):,}")
        c.drawRightString(520, y, f"${row['Total Sales ($)']:.2f}")
        y -= 18
        if y < 60:
            c.showPage()
            y = 800

    if include_chart and not summary.empty:
        chart_file = "temp_chart.png"
        try:
            chart = summary.set_index("Product")["Total Sales ($)"].sort_values(ascending=False).head(10)
            chart.plot(kind="bar", figsize=(6,4))
            plt.tight_layout()
            plt.savefig(chart_file, dpi=150)
            plt.close()
            c.showPage()
            # place image on new page
            c.drawImage(chart_file, 50, 300, width=500, preserveAspectRatio=True)
        finally:
            if os.path.exists(chart_file):
                os.remove(chart_file)

    c.save()

# --- MySQL storage ---
def store_to_mysql(df):
    try:
        conn = mysql.connector.connect(**DB_CONFIG)
        cursor = conn.cursor()
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS sales_log (
                id INT AUTO_INCREMENT PRIMARY KEY,
                product VARCHAR(255),
                units_sold INT,
                unit_price FLOAT,
                total_sales FLOAT,
                sale_date DATE
            )
        """)
        insert_sql = """
            INSERT INTO sales_log (product, units_sold, unit_price, total_sales, sale_date)
            VALUES (%s, %s, %s, %s, %s)
        """
        for _, row in df.iterrows():
            sale_date = None
            if pd.notna(row.get("Date")):
                if hasattr(row["Date"], "date"):
                    sale_date = row["Date"].date()
                else:
                    sale_date = row["Date"]
            cursor.execute(insert_sql, (
                str(row.get("Product")),
                int(row.get("Units Sold")),
                float(row.get("Unit Price ($)")),
                float(row.get("Total Sales ($)")),
                sale_date
            ))
        conn.commit()
        cursor.close()
        conn.close()
    except Exception as e:
        # do not block the flow; just report in console and let user continue
        print("MySQL insert failed:", e)

# --- GUI (layout from the design you liked) ---
root = tk.Tk()
root.title("Sales Report Generator")
root.geometry("450x260")
root.resizable(False, False)

file_path_var = tk.StringVar()
include_chart_var = tk.BooleanVar(value=True)
export_format = tk.StringVar(value="word")

# UI widgets (same layout style)
tk.Label(root, text="📂 Select Sales Data File", font=("Segoe UI", 10, "bold")).pack(pady=10)
file_entry = tk.Entry(root, textvariable=file_path_var, width=58)
file_entry.pack()

def on_browse():
    p = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xls"), ("CSV files", "*.csv"), ("All files", "*.*")])
    if p:
        file_path_var.set(p)

tk.Button(root, text="Browse", command=on_browse).pack(pady=6)

tk.Checkbutton(root, text="Include Chart in Report", variable=include_chart_var).pack()

# Export format radio buttons (placed below chart checkbox to match requested design)
fmt_frame = tk.Frame(root)
tk.Label(fmt_frame, text="Export Format:", font=("Segoe UI", 9, "bold")).pack(side=tk.LEFT, padx=(0,8))
tk.Radiobutton(fmt_frame, text="Word (.docx)", variable=export_format, value="word").pack(side=tk.LEFT)
tk.Radiobutton(fmt_frame, text="PDF (.pdf)", variable=export_format, value="pdf").pack(side=tk.LEFT, padx=(6,0))
fmt_frame.pack(pady=(8,0))

# Generate button (green)
def on_generate():
    path = file_path_var.get().strip()
    if not path:
        messagebox.showwarning("No file selected", "Please select a sales data file first.")
        return

    df = load_data(path)
    if df is None or df.empty:
        messagebox.showerror("No data", "No valid data could be loaded from the file.")
        return

    # prepare summary (top 10 by total sales)
    summary = df.groupby("Product", as_index=False).agg({
        "Units Sold": "sum",
        "Total Sales ($)": "sum"
    }).sort_values(by="Total Sales ($)", ascending=False).reset_index(drop=True)

    # filename
    today = datetime.now().strftime("%d_%b_%Y")
    base = f"sales_report_{today}"
    if export_format.get() == "word":
        out = base + ".docx"
        generate_word_report(summary.head(10), out, include_chart_var.get())
    else:
        out = base + ".pdf"
        generate_pdf_report(summary.head(10), out, include_chart_var.get())

    # store raw df to mysql (non-blocking if fails)
    store_to_mysql(df)

    messagebox.showinfo("Success", f"✅ Report saved as:\n{out}\n✅ Data logged to MySQL (if configured).")
    root.destroy()  # auto-close on success

tk.Button(root, text="Generate Report", command=on_generate, bg="#4CAF50", fg="white", font=("Segoe UI", 10, "bold")).pack(pady=14)

root.mainloop()
